"""
Builds the knowledge-base Word volumes from the Markdown drafts, using the Oimbra template
(fuentes/plantilla/Nomenclatura BD v1.0.docx) restyled with the Cotown Group brand.

Usage (from knowledge/): python scripts/gen_docx.py [version] [date]
Output: salida/*.docx (the table of contents is filled in by scripts/update_docx.ps1, which uses Word)
"""
import copy, io, os, re, sys, datetime, urllib.parse
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, 'fuentes', 'plantilla', 'Nomenclatura BD v1.0.docx')
LOGO_CG = os.path.join(ROOT, 'fuentes', 'plantilla', 'cotowngroup-isologo.png')
LOGO_OIMBRA = os.path.join(ROOT, 'fuentes', 'plantilla', 'oimbra-logo-light-crop.png')
DRAFTS = os.path.join(ROOT, 'borradores')
OUT = os.path.join(ROOT, 'salida')

VERSION = sys.argv[1] if len(sys.argv) > 1 else '0.1'
DATE = datetime.date.fromisoformat(sys.argv[2]) if len(sys.argv) > 2 else datetime.date.today()
MONTHS = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']

# Cotown Group brand
TEAL, PINK, BLACK = '08BDBD', 'FF4870', '000000'
OLD_GREEN, OLD_LINE = '47D45A', '414573'
TEXT_WIDTH_CM = 16.0

VOLUMES = [
  {
    'file': 'Cotown - Base de conocimiento - Manual técnico',
    'subtitle': 'Manual técnico',
    'chapters': ['01-vision-general.md', 'IT-01-arquitectura.md', 'IT-02-modelo-datos.md', 'IT-03-servicio-integracion.md',
                 'IT-04-procesos-automaticos.md', 'IT-05-triggers.md', 'IT-06-frontends.md', 'IT-07-webs.md', 'IT-08-data-warehouse.md', 'IT-09-integraciones.md', 'IT-10-infraestructura.md', 'IT-11-usuarios-roles.md', 'IT-12-instalacion.md', 'IT-13-deuda-tecnica.md', 'IT-A1-diccionario-datos.md', 'IT-A2-permisos.md'],
  },
  {
    'file': 'Cotown - Base de conocimiento - Manual funcional',
    'subtitle': 'Manual funcional',
    'chapters': ['01-vision-general.md', 'OP-01-inventario-precios.md', 'OP-02-captacion.md', 'OP-03-reservas-b2c.md', 'OP-04-reservas-grupo.md', 'OP-05-contratos-garantias.md', 'OP-06-facturacion-cobros.md', 'OP-07-operacion-diaria.md', 'OP-08-comunicacion-residente.md', 'OP-09-contratos-lau.md', 'OP-10-propietarios.md', 'OP-11-informes-analitica.md', 'OP-A1-calendario-procesos.md'],
  },
]

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


# ------------------------------------------------------------------ template restyling

def recolor(element):
  """Replace the template green and line colours by the brand teal in any XML subtree."""
  for el in element.iter():
    for k, v in list(el.attrib.items()):
      if v in (OLD_GREEN, OLD_GREEN.lower(), OLD_LINE):
        el.set(k, TEAL)


def picture_size(path, width_cm=None, height_cm=None):
  w, h = Image.open(path).size
  if width_cm:
    return Cm(width_cm), Cm(width_cm * h / w)
  return Cm(height_cm * w / h), Cm(height_cm)


def set_header(section):
  hdr = section.header
  p = next(p for p in hdr.paragraphs if p._p.findall('.//' + qn('w:drawing')))
  for r in list(p.runs):
    if r._r.findall('.//' + qn('wp:inline')):
      p._p.remove(r._r)
  p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
  w, h = picture_size(LOGO_CG, height_cm=0.6)
  p.add_run().add_picture(LOGO_CG, width=w, height=h)
  p.add_run('\t')
  w, h = picture_size(LOGO_OIMBRA, height_cm=0.55)
  p.add_run().add_picture(LOGO_OIMBRA, width=w, height=h)
  recolor(hdr._element)


def set_footer(section, subtitle):
  ftr = section.footer
  recolor(ftr._element)
  texts = [t for t in ftr._element.iter(qn('w:t'))]
  # Template sequence: "Oimbra ", "Engine", " ", "–", " ", "Modelos de datos - ", "Nomenclatura", "P", "á", "g", "ina", ...
  new = ['Cotown', '', ' ', '–', ' ', 'Base de conocimiento – ', subtitle]
  for t, v in zip(texts, new):
    t.text = v
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')


def set_cover(doc, body_els, subtitle):
  # Cover logo (paragraph with the big picture before the title)
  pic_par = next(el for el in body_els[:15] if el.tag == qn('w:p') and el.findall('.//' + qn('wp:inline')))
  par = docx.text.paragraph.Paragraph(pic_par, doc)
  for r in list(par.runs):
    par._p.remove(r._r)
  w, h = picture_size(LOGO_CG, width_cm=9.5)
  par.add_run().add_picture(LOGO_CG, width=w, height=h)
  # Title and subtitle
  titles = [el for el in body_els[:20] if el.tag == qn('w:p') and docx.text.paragraph.Paragraph(el, doc).style.name == 'Portada 2'
            and docx.text.paragraph.Paragraph(el, doc).text.strip()]
  set_text(docx.text.paragraph.Paragraph(titles[0], doc), 'Base de conocimiento')
  set_text(docx.text.paragraph.Paragraph(titles[1], doc), subtitle)
  # Version and date
  for el in body_els[:30]:
    if el.tag != qn('w:p'):
      continue
    p = docx.text.paragraph.Paragraph(el, doc)
    if p.text.startswith('Versión'):
      set_text(p, f'Versión {VERSION}')
    elif re.match(r'\d+ de \w+ de \d{4}', p.text):
      set_text(p, f'{DATE.day} de {MONTHS[DATE.month - 1]} de {DATE.year}')
      # "Elaborado por" + small Oimbra logo
      q = OxmlElement('w:p'); el.addnext(q)
      qp = docx.text.paragraph.Paragraph(q, doc); qp.style = doc.styles['Párrafo']
      qp.paragraph_format.space_before = Pt(18)
      r = qp.add_run('Elaborado por  '); r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string('404040')
      w, h = picture_size(LOGO_OIMBRA, height_cm=0.7)
      qp.add_run().add_picture(LOGO_OIMBRA, width=w, height=h)
      break


def set_text(p, text):
  runs = p.runs
  runs[0].text = text
  for r in runs[1:]:
    r.text = ''


def set_versions(doc):
  t = doc.tables[0]
  row = t.rows[1].cells
  for c, v in zip(row, [VERSION, DATE.strftime('%d/%m/%Y'), 'Oimbra', 'Borrador para revisión']):
    set_text(c.paragraphs[0], v) if c.paragraphs[0].runs else c.paragraphs[0].add_run(v)
  for extra in t.rows[2:]:
    t._tbl.remove(extra._tr)
  for c in t.rows[0].cells:
    shade(c, 'D9F4F4')


def clear_body_after_toc(doc):
  """Remove the template TOC and body (keeps the "Índice" title and the final sectPr) and insert a fresh TOC field."""
  body = doc.element.body
  els = list(body.iterchildren())
  start = next(i for i, el in enumerate(els)
               if el.tag == qn('w:p') and docx.text.paragraph.Paragraph(el, doc).text.strip() == 'Índice')
  for el in els[start + 1:]:
    if el.tag != qn('w:sectPr'):
      body.remove(el)
  p = doc.add_paragraph(style='Párrafo')
  def fld(kind):
    r = OxmlElement('w:r'); f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), kind); r.append(f); return r
  p._p.append(fld('begin'))
  r = OxmlElement('w:r'); t = OxmlElement('w:instrText'); t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
  t.text = ' TOC \\o "1-3" \\h \\z \\u '; r.append(t); p._p.append(r)
  p._p.append(fld('separate'))
  r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = 'Actualice el índice (F9).'; r.append(t); p._p.append(r)
  p._p.append(fld('end'))


def trim_cover(doc, body_els, n=3):
  """Remove n empty paragraphs between the subtitle and the version line, to make room for "Elaborado por"."""
  idx = next(i for i, el in enumerate(body_els) if el.tag == qn('w:p') and docx.text.paragraph.Paragraph(el, doc).text.startswith('Versión'))
  removed = 0
  for el in reversed(body_els[:idx]):
    if removed == n:
      break
    if el.tag == qn('w:p') and not docx.text.paragraph.Paragraph(el, doc).text.strip() and not el.findall('.//' + qn('w:drawing')):
      el.getparent().remove(el); removed += 1


# ------------------------------------------------------------------ markdown rendering

INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def add_inline(p, text, size=None, bold=False):
  text = re.sub(r'\s*\*\*\[PREGUNTA [^\]]+\]\*\*', '', text)
  text = re.sub(r'<br\s*/?>', ' ', text)
  for part in INLINE.split(text):
    if not part:
      continue
    if part.startswith('**') and part.endswith('**'):
      inner = part[2:-2]
      if inner.startswith('`') and inner.endswith('`'):
        r = p.add_run(inner[1:-1]); r.bold = True; r.font.name = 'Consolas'; r.font.size = Pt(8.5)
      else:
        r = p.add_run(inner); r.bold = True
    elif part.startswith('`') and part.endswith('`'):
      r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
      r.font.color.rgb = RGBColor.from_string('404040')
    elif part.startswith('[') and '](' in part:
      r = p.add_run(re.match(r'\[([^\]]+)\]', part).group(1))
    elif part.startswith('*') and part.endswith('*') and len(part) > 2:
      r = p.add_run(part[1:-1]); r.italic = True
    else:
      r = p.add_run(part)
    if size:
      r.font.size = size
    if bold:
      r.bold = True


def shade(cell, fill):
  tcPr = cell._tc.get_or_add_tcPr()
  for s in tcPr.findall(qn('w:shd')):
    tcPr.remove(s)
  shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
  tcPr.append(shd)


def add_table(doc, rows):
  header, data = rows[0], rows[1:]
  ncol = len(header)
  t = doc.add_table(rows=0, cols=ncol)
  t.style = doc.styles['Table Grid']
  # Column widths proportional to content length (bounded), summing to the text width
  lens = []
  for c in range(ncol):
    m = max(len(re.sub(r'[*`]', '', r[c])) if c < len(r) else 0 for r in rows)
    lens.append(min(max(m, 6), 60))
  total = sum(lens)
  widths = [Cm(TEXT_WIDTH_CM * l / total) for l in lens]
  small = Pt(8.5) if ncol >= 4 or len(data) > 15 else Pt(9.5)
  for ri, r in enumerate(rows):
    cells = t.add_row().cells
    for c in range(ncol):
      cell = cells[c]
      cell.width = widths[c]
      p = cell.paragraphs[0]; p.style = doc.styles['Párrafo']
      p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
      p.alignment = WD_ALIGN_PARAGRAPH.LEFT
      add_inline(p, r[c] if c < len(r) else '', size=small, bold=(ri == 0))
      if ri == 0:
        shade(cell, 'D9F4F4')
  # Repeat header row on each page
  trPr = t.rows[0]._tr.get_or_add_trPr()
  h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
  doc.add_paragraph(style='Párrafo').paragraph_format.space_after = Pt(0)


def add_image(doc, path, alt):
  if not os.path.exists(path):
    p = doc.add_paragraph(style='Párrafo'); add_inline(p, f'*[Falta la imagen: {alt}]*'); return
  w, h = Image.open(path).size
  width = TEXT_WIDTH_CM
  if width * h / w > 21:
    width = 21 * w / h
  p = doc.add_paragraph(style='Párrafo'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  p.add_run().add_picture(path, width=Cm(width))
  c = doc.add_paragraph(style='Párrafo'); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
  r = c.add_run(alt); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string('404040')


def strip_number(title):
  title = re.sub(r'^(IT|OP)-\d+\.\s*', '', title)
  title = re.sub(r'^\d+(\.\d+)*\.?\s+', '', title)
  return title.strip()


def render_markdown(doc, path):
  lines = open(path, encoding='utf-8').read().splitlines()
  base = os.path.dirname(path)
  annex = re.match(r'^(IT|OP)-A\d', os.path.basename(path)) is not None
  i = 0
  para = []

  def flush():
    if para:
      p = doc.add_paragraph(style='Párrafo'); add_inline(p, ' '.join(s.strip() for s in para)); para.clear()

  while i < len(lines):
    line = lines[i]
    s = line.strip()
    if not s:
      flush(); i += 1; continue
    # Draft metadata lines
    if s.startswith('> ') and re.search(r'Borrador v|Capítulo común|Manual técnico ·|Manual funcional ·|Generado automáticamente', s):
      flush(); i += 1; continue
    m = re.match(r'^(#{1,4})\s+(.*)$', s)
    if m:
      flush()
      level, title = len(m.group(1)), m.group(2)
      if level == 4:
        p = doc.add_paragraph(style='Párrafo'); add_inline(p, title, bold=True)
      else:
        # Annex detail headings keep the look and numbering of Heading 3 but stay out of the table of contents
        style = ANNEX_H3 if (level == 3 and annex) else f'Heading {level}'
        p = doc.add_paragraph(style=style); add_inline(p, strip_number(title))
      i += 1; continue
    m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', s)
    if m:
      flush()
      add_image(doc, os.path.normpath(os.path.join(base, urllib.parse.unquote(m.group(2)))), m.group(1))
      i += 1; continue
    if s.startswith('```'):
      flush(); i += 1
      while i < len(lines) and not lines[i].strip().startswith('```'):
        p = doc.add_paragraph(style='Badge'); p.add_run(lines[i].replace('\t', '  ') or ' ')
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        i += 1
      i += 1; continue
    if s.startswith('|'):
      flush()
      rows = []
      while i < len(lines) and lines[i].strip().startswith('|'):
        row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
        if not all(re.match(r'^:?-{2,}:?$', c) for c in row if c):
          rows.append(row)
        i += 1
      add_table(doc, rows)
      continue
    m = re.match(r'^(\s*)[-*]\s+(.*)$', line)
    if m:
      flush()
      p = doc.add_paragraph(style='Bullets')
      if len(m.group(1)) >= 2:
        numPr = p._p.get_or_add_pPr().get_or_add_numPr()
        numPr.get_or_add_ilvl().val = 1
        numPr.get_or_add_numId().val = 4
      add_inline(p, m.group(2)); i += 1; continue
    m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
    if m:
      flush()
      p = doc.add_paragraph(style='Párrafo')
      p.paragraph_format.left_indent = Cm(0.9 + 0.6 * (len(m.group(1)) // 2)); p.paragraph_format.first_line_indent = Cm(-0.5)
      p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(4)
      add_inline(p, f'{m.group(2)}. {m.group(3)}'); i += 1; continue
    if s.startswith('>'):
      flush()
      p = doc.add_paragraph(style='Párrafo'); add_inline(p, f'*{s.lstrip("> ").strip()}*'); i += 1; continue
    para.append(s); i += 1
  flush()


# ------------------------------------------------------------------ build

ANNEX_H3 = 'Título 3 anexo'


def add_annex_style(doc):
  """Clone of Heading 3 with outline level 'body text', so it is numbered but not listed in the TOC."""
  styles = doc.styles.element
  h3 = next(s for s in styles.findall(qn('w:style')) if s.get(qn('w:styleId')) == doc.styles['Heading 3'].style_id)
  s = copy.deepcopy(h3)
  s.set(qn('w:styleId'), 'Titulo3Anexo')
  s.remove(s.find(qn('w:name'))); n = OxmlElement('w:name'); n.set(qn('w:val'), ANNEX_H3); s.insert(0, n)
  for tag in ('w:link', 'w:next'):
    e = s.find(qn(tag))
    if e is not None:
      s.remove(e)
  s.set(qn('w:customStyle'), '1')
  ppr = s.find(qn('w:pPr'))
  ppr.find(qn('w:outlineLvl')).set(qn('w:val'), '9')
  styles.append(s)


def build(volume):
  doc = docx.Document(TEMPLATE)
  recolor(doc.styles.element)
  add_annex_style(doc)
  body_els = list(doc.element.body.iterchildren())
  recolor(doc.element.body)
  set_cover(doc, body_els, volume['subtitle'])
  trim_cover(doc, body_els)
  set_versions(doc)
  clear_body_after_toc(doc)
  for s in doc.sections:
    set_header(s)
    set_footer(s, volume['subtitle'])
  for ch in volume['chapters']:
    render_markdown(doc, os.path.join(DRAFTS, ch))
  # Word recalculates the TOC and fields when the document is opened
  settings = doc.settings.element
  upd = settings.find(qn('w:updateFields'))
  if upd is None:
    upd = OxmlElement('w:updateFields'); settings.append(upd)
  upd.set(qn('w:val'), 'true')
  props = doc.core_properties
  props.title = f"Cotown – Base de conocimiento – {volume['subtitle']}"
  props.author = 'Oimbra'
  props.subject = 'Base de conocimiento'
  os.makedirs(OUT, exist_ok=True)
  path = os.path.join(OUT, f"{volume['file']} v{VERSION}.docx")
  doc.save(path)
  print('ok', path)


if __name__ == '__main__':
  for v in VOLUMES:
    build(v)
